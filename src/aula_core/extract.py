"""Text extraction from course material, one entry per page / slide / section, and a PDF page
as an image (for a scanned page the text extraction cannot read)."""

from __future__ import annotations

import io
import logging
from pathlib import Path

log = logging.getLogger(__name__)
# pypdf logs a warning per odd font or repaired object, each dumping the font dictionary: a first
# download filled bot.log at ~1.7 MB a minute. The text comes out the same.
logging.getLogger("pypdf").setLevel(logging.ERROR)

# What a "page" means for each format, used when citing.
UNIT = {"pdf": "página", "pptx": "diapositiva", "docx": "sección", "html": "sección"}

DOCX_SECTION_CHARS = 1800
PAGE_IMAGE_EDGE = 1600  # pixels on the long side: readable, and still well under the model's image limits


def extract(path: Path) -> list[tuple[int, str]]:
    ext = path.suffix.lower().lstrip(".")
    if ext == "pdf":
        return _pdf(path)
    if ext == "pptx":
        return _pptx(path)
    if ext == "docx":
        return _docx(path)
    if ext in ("html", "htm"):
        return _html(path)
    raise ValueError(f"Formato no soportado: {path.suffix}")


def layout_text(path: Path, max_pages: int = 12) -> str:
    """A PDF's first pages keeping the columns apart (a syllabus puts BÁSICA and COMPLEMENTARIA side by side)."""
    from pypdf import PdfReader

    parts = []
    for page in PdfReader(str(path)).pages[:max_pages]:
        try:
            parts.append(page.extract_text(extraction_mode="layout") or "")
        except Exception as exc:
            log.warning("%s: %s", path.name, exc)
    return "\n".join(parts)


def render_page(path: Path, number: int) -> bytes:
    """Page `number` (from 1) of a PDF as a JPEG."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        if not 1 <= number <= len(pdf):
            raise ValueError(f"{path.name} tiene {len(pdf)} páginas; no hay página {number}.")
        page = pdf[number - 1]
        width, height = page.get_size()
        image = page.render(scale=PAGE_IMAGE_EDGE / max(width, height, 1)).to_pil().convert("RGB")
        out = io.BytesIO()
        image.save(out, "JPEG", quality=80)
        return out.getvalue()
    finally:
        pdf.close()


def _pdf(path: Path) -> list[tuple[int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # a single broken page should not lose the whole file
            log.warning("%s p.%d: %s", path.name, number, exc)
            text = ""
        pages.append((number, text.strip()))
    return pages


def _pptx(path: Path) -> list[tuple[int, str]]:
    from pptx import Presentation

    slides = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notas: {notes}")
        slides.append((number, "\n".join(p for p in parts if p.strip()).strip()))
    return slides


def _docx(path: Path) -> list[tuple[int, str]]:
    from docx import Document

    doc = Document(str(path))
    blocks = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    sections: list[tuple[int, str]] = []
    current: list[str] = []
    size = 0
    for block in blocks:
        current.append(block)
        size += len(block)
        if size >= DOCX_SECTION_CHARS:
            sections.append((len(sections) + 1, "\n".join(current)))
            current, size = [], 0
    if current:
        sections.append((len(sections) + 1, "\n".join(current)))
    return sections


def _html(path: Path) -> list[tuple[int, str]]:
    from aula_core.catalog import html_to_text

    text = html_to_text(path.read_text(encoding="utf-8", errors="replace"))
    sections, current = [], []
    for block in text.split("\n\n"):
        current.append(block)
        if sum(map(len, current)) >= DOCX_SECTION_CHARS:
            sections.append((len(sections) + 1, "\n\n".join(current)))
            current = []
    if current:
        sections.append((len(sections) + 1, "\n\n".join(current)))
    return sections
