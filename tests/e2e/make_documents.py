"""Regenerates the binary course-material fixtures (committed under fixtures/files/).

    uv run python tests/e2e/make_documents.py

The PDFs with text are written by hand (standard Helvetica font, WinAnsi encoding) so the
fixtures need no PDF library; the PPTX uses python-pptx, and the scanned reading is page images
saved as a PDF with Pillow (no text at all, like a photocopy run through a scanner).

- capitulo-3-derivadas.pdf, semana-2-cinematica.pptx: course material with text;
- silabo-matg1049.pdf: an ESPOL «CONTENIDO DE ASIGNATURA» with BÁSICA and COMPLEMENTARIA side by side;
- purcell-calculo.pdf: the main book of Cálculo (Spanish), which the captain sends the bot;
- serway-physics.pdf: the main book of Física (English), which the captain drops in libros/FISG1002/;
- lectura-vectores-escaneada.pdf: a scanned reading of Física (a formula OCR misreads);
- ejercicios-derivadas-escaneados.pdf: a scanned exercise sheet the captain sends Cálculo (prose OCR reads well);
- politicas-del-curso.pdf, guia-laboratorio-1.pdf: what Física's Google Doc and Drive file hand over (FakeWeb).
"""

from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).parent / "fixtures" / "files"

PDF_PAGES = [
    [
        "Calculo Diferencial - Capitulo 3: Derivadas",
        "3.1 Definicion de la derivada",
        "La derivada de f en x es el limite del cociente incremental",
        "(f(x+h) - f(x)) / h cuando h tiende a cero.",
        "Geometricamente es la pendiente de la recta tangente a la curva.",
    ],
    [
        "3.2 La regla de la cadena",
        "La regla de la cadena permite derivar una composicion de funciones.",
        "Si y = f(g(x)), entonces y' = f'(g(x)) * g'(x).",
        "Ejemplo: la derivada de sen(x^2) es cos(x^2) * 2x.",
        "Se aplica de afuera hacia adentro, capa por capa.",
    ],
    [
        "3.3 Aplicaciones: optimizacion",
        "Para hallar maximos y minimos se iguala la derivada a cero.",
        "Los puntos criticos se clasifican con el criterio de la segunda derivada.",
        "Ejercicio 4: maximizar el area de un rectangulo de perimetro 20 m.",
    ],
]

# Accented words are added per line so the extractor must decode WinAnsi bytes.
ACCENTS = {
    "Calculo": "Cálculo", "Capitulo": "Capítulo", "Definicion": "Definición", "limite": "límite",
    "Geometricamente": "Geométricamente", "composicion": "composición", "optimizacion": "optimización",
    "maximos": "máximos", "minimos": "mínimos", "criticos": "críticos", "area": "área",
    "perimetro": "perímetro",
}


def _accented(line: str) -> str:
    for plain, accented in ACCENTS.items():
        line = line.replace(plain, accented)
    return line


def _pdf_string(text: str) -> str:
    out = []
    for byte in text.encode("cp1252"):
        ch = chr(byte)
        if ch in "()\\":
            out.append("\\" + ch)
        elif 32 <= byte < 127:
            out.append(ch)
        else:
            out.append(f"\\{byte:03o}")
    return "(" + "".join(out) + ")"


def make_pdf(path: Path) -> None:
    objects: list[bytes] = []
    n_pages = len(PDF_PAGES)
    # 1 catalog, 2 pages, 3 font, then (page, content) pairs
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pages))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for i, lines in enumerate(PDF_PAGES):
        stream_lines = ["BT", "/F1 12 Tf", "72 740 Td", "16 TL"]
        for line in lines:
            stream_lines.append(f"{_pdf_string(_accented(line))} Tj T*")
        stream_lines.append("ET")
        stream = "\n".join(stream_lines).encode("latin-1")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    body = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        body += f"{offset:010d} 00000 n \n".encode()
    body += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(body))


def write_pdf(path: Path, pages: list[list[tuple[int, int, str]]]) -> None:
    """A text PDF from (x, y, text) runs per page, in points from the bottom left."""
    objects: list[bytes] = []
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for i, runs in enumerate(pages):
        stream = "\n".join(["BT", "/F1 10 Tf", *(f"1 0 0 1 {x} {y} Tm {_pdf_string(text)} Tj" for x, y, text in runs),
                             "ET"]).encode("latin-1")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    body = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        body += f"{offset:010d} 00000 n \n".encode()
    body += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(body))


def _lines(lines: list[str], top: int = 740, x: int = 72, step: int = 16) -> list[tuple[int, int, str]]:
    return [(x, top - i * step, line) for i, line in enumerate(lines)]


def make_syllabus(path: Path) -> None:
    left, right = 72, 330
    head = _lines([
        "ESCUELA SUPERIOR POLITÉCNICA DEL LITORAL",
        "CONTENIDO DE ASIGNATURA",
        "MATG1049 - CÁLCULO DE UNA VARIABLE",
        "A. INFORMACIÓN GENERAL",
        "Créditos: 4. Horas de clase por semana: 4 de teoría y 2 de práctica.",
        "B. RESULTADOS DE APRENDIZAJE",
        "Al terminar el curso el estudiante resuelve problemas de la ingeniería con las herramientas",
        "del cálculo diferencial e integral de una variable real.",
        "C. CONTENIDOS",
        "Unidad 1: Límites y continuidad de funciones.",
        "Unidad 2: La derivada y sus propiedades.",
        "Unidad 3: Aplicaciones de la derivada.",
        "Unidad 4: La integral definida.",
        "I. BIBLIOGRAFÍA",
    ])
    y = head[-1][1] - 22
    table = [(left, y, "BÁSICA"), (right, y, "COMPLEMENTARIA")]
    rows = [
        ("1. Purcell, E., Varberg, D. y Rigdon, S.", "1. Stewart, J. (2012). Cálculo de una"),
        ("(2007). Cálculo (9a ed.). Pearson.", "variable: trascendentes tempranas (7a ed.)."),
        ("", "Cengage Learning."),
        ("", "2. Larson, R. y Edwards, B. (2010). Cálculo"),
        ("", "(9a ed.). McGraw-Hill."),
    ]
    for i, (a, b) in enumerate(rows, start=1):
        row_y = y - i * 14
        table += [(left, row_y, a)] if a else []
        table += [(right, row_y, b)] if b else []
    tail = _lines(["J. EVALUACIÓN", "Primer parcial 35%, segundo parcial 35%, deberes y lecciones 30%."],
                  top=y - (len(rows) + 2) * 14)
    write_pdf(path, [head + table + tail])


def make_purcell(path: Path) -> None:
    write_pdf(path, [
        _lines(["Purcell, Varberg y Rigdon - Cálculo, 9a edición", "Capítulo 2: La derivada",
                "2.1 Dos problemas con el mismo tema: la recta tangente y la velocidad instantánea.",
                "La derivada de una función es otra función que da la pendiente de la tangente."]),
        _lines(["2.5 La regla de la cadena",
                "Teorema A (Regla de la cadena). Si y = f(u) y u = g(x), entonces la derivada de la",
                "composición es Dx y = Du y · Dx u, es decir (f o g)'(x) = f'(g(x)) · g'(x).",
                "Se deriva la función de afuera evaluada en la de adentro y se multiplica por la",
                "derivada de la función de adentro. Ejemplo: si y = (2x^2 - 4x + 1)^60, entonces",
                "Dx y = 60 (2x^2 - 4x + 1)^59 · (4x - 4)."]),
    ])


def make_serway(path: Path) -> None:
    write_pdf(path, [
        _lines(["Serway and Jewett - Physics for Scientists and Engineers", "Chapter 4: Motion in Two Dimensions",
                "4.1 The position, velocity, and acceleration vectors of a particle moving in a plane."]),
        _lines(["4.3 Projectile Motion",
                "A projectile moves with constant horizontal velocity and constant downward acceleration g.",
                "Its path is a parabola. The horizontal range of a projectile launched with speed vi at",
                "angle theta is R = vi^2 sin(2 theta) / g, so the maximum range is reached when theta is 45",
                "degrees. The maximum height is h = vi^2 sin^2(theta) / (2g)."]),
    ])


def make_policies(path: Path) -> None:
    write_pdf(path, [
        _lines(["Física I - Políticas del curso (II PAO 2026)",
                "Asistencia: la asistencia mínima para aprobar es del 70 % de las clases.",
                "Los deberes atrasados se reciben hasta 24 horas después, con 20 % menos.",
                "El celular va en silencio durante la clase."]),
        _lines(["Evaluación: primer parcial 30 %, segundo parcial 30 %, laboratorio 20 %, deberes 20 %.",
                "La nota del laboratorio sale de los informes de cada práctica."]),
    ])


def make_lab_guide(path: Path) -> None:
    write_pdf(path, [
        _lines(["Física I - Laboratorio 1: movimiento en un plano inclinado",
                "Materiales: riel de aire, carrito, cronómetro y cinta métrica.",
                "Procedimiento: suelte el carrito desde 1 m de altura y mida el tiempo de bajada cinco veces.",
                "Informe: tabla de tiempos, aceleración media y comparación con g sen(theta)."]),
    ])


SCAN_PAGES = [["Lectura 1: suma de vectores", "Regla del paralelogramo:", "R = A + B, |R|^2 = A^2 + B^2 + 2AB cos(t)"],
              ["Componentes de un vector", "Ax = A cos(t)   Ay = A sen(t)", "A = Ax i + Ay j"]]
EXERCISE_PAGES = [["Hoja de ejercicios 2: derivadas", "Resuelve cada ejercicio y justifica cada paso.",
                   "Ejercicio 4: deriva el producto de x al cubo por el seno de x.",
                   "Ejercicio 5: con la regla de la cadena, deriva el coseno de x^2."]]


def make_scan(path: Path, sheets: list[list[str]] = SCAN_PAGES, title: str = "Lectura 1") -> None:
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.load_default(size=34)
    pages = []
    for lines in sheets:
        page = Image.new("L", (1240, 1754), 250)
        draw = ImageDraw.Draw(page)
        for i, line in enumerate(lines):
            draw.text((120, 160 + i * 90), line, fill=20, font=font)
        pages.append(page.convert("RGB"))
    moment = __import__("datetime").datetime(2026, 9, 1).timetuple()  # fixed, so the file is byte-stable
    pages[0].save(path, "PDF", resolution=150, save_all=True, append_images=pages[1:], creationDate=moment,
                  modDate=moment, producer="scanner", title=title)


def make_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slides = [
        ("Física I - Semana 2: Cinemática", "Movimiento en una y dos dimensiones"),
        ("Movimiento rectilíneo uniforme", "La velocidad es constante: x = x0 + v·t"),
        ("Movimiento parabólico", "Un proyectil combina MRU en el eje x y caída libre en el eje y.\n"
                                  "El alcance máximo se logra con un ángulo de 45°."),
    ]
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.core_properties.author = "ESPOL"
    # Fixed timestamps keep the file byte-stable across regenerations.
    import datetime as _dt
    prs.core_properties.created = _dt.datetime(2026, 9, 1)
    prs.core_properties.modified = _dt.datetime(2026, 9, 1)
    prs.core_properties.last_modified_by = "ESPOL"
    prs.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    make_pdf(OUT / "capitulo-3-derivadas.pdf")
    make_pptx(OUT / "semana-2-cinematica.pptx")
    make_syllabus(OUT / "silabo-matg1049.pdf")
    make_purcell(OUT / "purcell-calculo.pdf")
    make_serway(OUT / "serway-physics.pdf")
    make_scan(OUT / "lectura-vectores-escaneada.pdf")
    make_scan(OUT / "ejercicios-derivadas-escaneados.pdf", EXERCISE_PAGES, "Hoja de ejercicios 2")
    make_policies(OUT / "politicas-del-curso.pdf")
    make_lab_guide(OUT / "guia-laboratorio-1.pdf")
    print(f"Fixtures escritos en {OUT}")
