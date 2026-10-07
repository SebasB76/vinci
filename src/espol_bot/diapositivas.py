"""A slide deck (.pptx) a bot builds from an outline it wrote, for a class presentation.

The file goes to the profile's `cache/documents/`, one of the folders Hermes delivers from: a reply that
mentions its path reaches the chat as a file (Telegram or WhatsApp), and the path itself is cut from the text.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

MAX_SLIDES = 30
MAX_POINTS = 8
TITLE_LAYOUT, CONTENT_LAYOUT = 0, 1


class SlidesError(ValueError):
    """An outline the deck cannot be built from."""


def _slug(title: str) -> str:
    plain = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "-", plain).strip("-").lower()[:60] or "diapositivas"


def _text(value, name: str, limit: int) -> str:
    text = str(value or "").strip()
    if not text:
        raise SlidesError(f"falta «{name}»")
    return text[:limit]


def _widen(slide, width) -> None:
    # The default template's boxes are laid out for 4:3; stretch them across the 16:9 page.
    for shape in slide.placeholders:
        shape.left, shape.width = Inches(0.7), width - Inches(1.4)


def build(folder: Path, title: str, slides: list[dict], subtitle: str = "", now: datetime | None = None) -> Path:
    title = _text(title, "titulo", 120)
    if not slides:
        raise SlidesError("faltan las diapositivas: manda al menos una con su título y sus puntos")
    if len(slides) > MAX_SLIDES:
        raise SlidesError(f"son {len(slides)} diapositivas; el máximo es {MAX_SLIDES}")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)  # 16:9, what a projector expects

    cover = deck.slides.add_slide(deck.slide_layouts[TITLE_LAYOUT])
    cover.shapes.title.text = title
    cover.placeholders[1].text = str(subtitle or "").strip()[:200]
    _widen(cover, deck.slide_width)

    for n, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            raise SlidesError(f"la diapositiva {n} no es un objeto con «titulo» y «puntos»")
        page = deck.slides.add_slide(deck.slide_layouts[CONTENT_LAYOUT])
        _widen(page, deck.slide_width)
        page.shapes.title.text = _text(slide.get("titulo"), f"titulo de la diapositiva {n}", 120)
        points = [str(p).strip() for p in slide.get("puntos") or [] if str(p).strip()][:MAX_POINTS]
        body = page.placeholders[1].text_frame
        body.clear()
        for i, point in enumerate(points):
            paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
            paragraph.text = point[:300]
            paragraph.font.size = Pt(24 if len(points) <= 5 else 20)
        notes = str(slide.get("notas") or "").strip()
        if notes:
            page.notes_slide.notes_text_frame.text = notes[:2000]

    folder.mkdir(parents=True, exist_ok=True)
    stamp = (now or datetime.now()).strftime("%Y%m%d-%H%M%S")
    path = folder / f"{_slug(title)}-{stamp}.pptx"
    deck.save(path)
    return path
